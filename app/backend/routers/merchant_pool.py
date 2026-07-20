import csv
import io
import json
import re
from datetime import datetime, timezone
from typing import Any, Optional
from urllib.parse import urlsplit

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile, status
from openpyxl import load_workbook
from pydantic import BaseModel, Field, ValidationError, field_validator
from sqlalchemy import func, or_, select
from sqlalchemy.ext.asyncio import AsyncSession

from core.database import get_db
from dependencies.auth import get_current_user
from models.customers import Customers
from models.merchant_pool import MerchantPool
from models.merchant_ai_analyses import MerchantAiAnalyses
from models.sales_leads import SalesLeads
from models.employees import Employees
from schemas.auth import UserResponse


router = APIRouter(prefix="/api/v1/merchant-pool", tags=["merchant-pool"])

ADMIN_ROLES = {"admin", "super_admin"}
POOL_ROLES = ADMIN_ROLES | {"sales_manager"}
ISOLATED_STATUSES = {"no_phone", "duplicate", "existing_customer", "closed"}
ARCHIVED_STATUS = "archived"
CLOSED_MARKERS = {"closed", "permanently_closed", "关闭", "已关闭", "停业", "暂停营业"}


def _role(user: UserResponse) -> str:
    return str(user.role or "").strip().lower()


def _employee_id(user: UserResponse) -> int:
    try:
        return int(user.id)
    except (TypeError, ValueError):
        raise HTTPException(status_code=403, detail="当前账号未关联有效员工")


def _ensure_pool_role(user: UserResponse) -> None:
    if _role(user) not in POOL_ROLES:
        raise HTTPException(status_code=403, detail="只有销售主管或管理员可以访问待清洗商家池")


def _ensure_admin_role(user: UserResponse) -> None:
    if _role(user) not in ADMIN_ROLES:
        raise HTTPException(status_code=403, detail="只有系统管理员可以清理商家池数据")


def _normalize_phone(value: Optional[str]) -> str:
    raw = re.sub(r"(?:ext\.?|extension|x|分机)\s*\d+\s*$", "", value or "", flags=re.I)
    digits = re.sub(r"[^0-9]", "", raw)
    # North American numbers are commonly imported both with and without +1.
    if len(digits) == 11 and digits.startswith("1"):
        return digits[1:]
    return digits


def _normalize_text(value: Optional[str]) -> str:
    return re.sub(r"[\W_]+", "", (value or "").casefold(), flags=re.UNICODE)


def _normalize_website(value: Optional[str]) -> str:
    website = (value or "").strip().casefold()
    url_match = re.search(r"https?://[^\s\]\)]+", website, re.I)
    if url_match:
        website = url_match.group().rstrip(".,;:")
    if not website:
        return ""
    parsed = urlsplit(website if "://" in website else f"https://{website}")
    host = (parsed.hostname or "").removeprefix("www.")
    path = re.sub(r"/+", "/", parsed.path or "").rstrip("/")
    return f"{host}{path}" if host else ""


def _is_closed(value: Optional[str]) -> bool:
    normalized = (value or "").strip().casefold()
    return normalized in CLOSED_MARKERS


class MerchantRecord(BaseModel):
    business_name: str = Field(min_length=1, max_length=200)
    contact_name: Optional[str] = None
    phone: Optional[str] = None
    industry: Optional[str] = None
    country: Optional[str] = None
    state: Optional[str] = None
    city: Optional[str] = None
    address: Optional[str] = None
    website: Optional[str] = None
    rating: Optional[float] = Field(default=None, ge=0, le=5)
    google_business_url: Optional[str] = None
    google_rating: Optional[float] = Field(default=None, ge=0, le=5)
    google_review_count: Optional[int] = Field(default=None, ge=0)
    yelp_url: Optional[str] = None
    yelp_rating: Optional[float] = Field(default=None, ge=0, le=5)
    yelp_review_count: Optional[int] = Field(default=None, ge=0)
    social_profiles: Optional[str] = None
    recent_negative_reviews: Optional[str] = None
    content_update_summary: Optional[str] = None
    content_last_updated_at: Optional[datetime] = None
    source_record_id: Optional[str] = None
    business_status: Optional[str] = None
    collected_at: Optional[datetime] = None

    @field_validator("business_name")
    @classmethod
    def strip_business_name(cls, value: str) -> str:
        return value.strip()


class MerchantImportRequest(BaseModel):
    data_source: str = Field(default="api", min_length=1, max_length=50)
    # Keep the original keys so later enrichment can inspect imported evidence.
    records: list[dict[str, Any]] = Field(min_length=1, max_length=2000)


class MerchantUpdate(BaseModel):
    business_name: Optional[str] = None
    contact_name: Optional[str] = None
    phone: Optional[str] = None
    industry: Optional[str] = None
    country: Optional[str] = None
    state: Optional[str] = None
    city: Optional[str] = None
    address: Optional[str] = None
    website: Optional[str] = None
    rating: Optional[float] = Field(default=None, ge=0, le=5)
    google_business_url: Optional[str] = None
    google_rating: Optional[float] = Field(default=None, ge=0, le=5)
    google_review_count: Optional[int] = Field(default=None, ge=0)
    yelp_url: Optional[str] = None
    yelp_rating: Optional[float] = Field(default=None, ge=0, le=5)
    yelp_review_count: Optional[int] = Field(default=None, ge=0)
    social_profiles: Optional[str] = None
    recent_negative_reviews: Optional[str] = None
    content_update_summary: Optional[str] = None
    content_last_updated_at: Optional[datetime] = None
    data_source: Optional[str] = None
    business_status: Optional[str] = None


class ConvertToLeadRequest(BaseModel):
    assigned_sales_id: Optional[int] = None


class BulkConvertToLeadRequest(BaseModel):
    merchant_ids: list[int] = Field(min_length=1, max_length=200)
    assigned_sales_id: int


class BulkMerchantIdsRequest(BaseModel):
    merchant_ids: list[int] = Field(min_length=1, max_length=200)


class BulkIndustryUpdateRequest(BulkMerchantIdsRequest):
    industry: str = Field(min_length=1, max_length=100)

    @field_validator("industry")
    @classmethod
    def strip_industry(cls, value: str) -> str:
        value = value.strip()
        if not value:
            raise ValueError("行业不能为空")
        return value


class MerchantResponse(BaseModel):
    id: int
    business_name: str
    contact_name: Optional[str] = None
    phone: Optional[str] = None
    industry: Optional[str] = None
    country: Optional[str] = None
    state: Optional[str] = None
    city: Optional[str] = None
    address: Optional[str] = None
    website: Optional[str] = None
    rating: Optional[float] = None
    google_business_url: Optional[str] = None
    google_rating: Optional[float] = None
    google_review_count: Optional[int] = None
    yelp_url: Optional[str] = None
    yelp_rating: Optional[float] = None
    yelp_review_count: Optional[int] = None
    social_profiles: Optional[str] = None
    recent_negative_reviews: Optional[str] = None
    content_update_summary: Optional[str] = None
    content_last_updated_at: Optional[datetime] = None
    data_source: str
    source_record_id: Optional[str] = None
    collected_at: datetime
    business_status: Optional[str] = None
    pool_status: str
    isolation_reason: Optional[str] = None
    duplicate_of_id: Optional[int] = None
    existing_customer_id: Optional[int] = None
    converted_lead_id: Optional[int] = None
    created_at: datetime
    updated_at: datetime

    class Config:
        from_attributes = True


class MerchantListResponse(BaseModel):
    items: list[MerchantResponse]
    total: int
    skip: int
    limit: int


CSV_HEADERS = {
    "商家名称": "business_name", "商家": "business_name", "商户名称": "business_name", "企业名称": "business_name", "店铺名称": "business_name", "名称": "business_name", "business_name": "business_name", "business name": "business_name", "name": "business_name", "title": "business_name",
    "联系人": "contact_name", "联系人姓名": "contact_name", "负责人": "contact_name", "contact_name": "contact_name", "contact name": "contact_name", "contact": "contact_name",
    "电话": "phone", "电话/网站": "phone", "电话网站": "phone", "联系电话": "phone", "电话号码": "phone", "手机": "phone", "手机号": "phone", "phone": "phone", "phone_number": "phone", "phone number": "phone", "telephone": "phone", "tel": "phone",
    "行业": "industry", "地区/行业": "industry", "地区行业": "industry", "类别": "industry", "分类": "industry", "商家类别": "industry", "industry": "industry", "category": "industry", "primary category": "industry",
    "地区": "region", "location": "region", "region": "region",
    "国家": "country", "国家地区": "country", "country": "country", "州": "state", "省": "state", "州省": "state", "state": "state", "province": "state",
    "城市": "city", "city": "city", "地址": "address", "详细地址": "address", "营业地址": "address", "address": "address", "full address": "address",
    "网站": "website", "官网": "website", "官网链接": "website", "网站链接": "website", "website": "website", "url": "website", "web site": "website", "domain": "website",
    "评分": "rating", "公开评分": "rating", "谷歌评分": "rating", "星级": "rating", "rating": "rating", "来源记录id": "source_record_id", "来源id": "source_record_id", "source_record_id": "source_record_id", "place id": "source_record_id",
    "营业状态": "business_status", "商家状态": "business_status", "状态": "business_status", "business_status": "business_status", "business status": "business_status", "status": "business_status",
    "采集时间": "collected_at", "抓取时间": "collected_at", "更新时间": "collected_at", "collected_at": "collected_at", "collected at": "collected_at",
    "google商家链接": "google_business_url", "google地图链接": "google_business_url", "google maps url": "google_business_url", "google_business_url": "google_business_url",
    "google评分": "google_rating", "Google评分": "google_rating", "google rating": "google_rating", "google_rating": "google_rating", "google评论数": "google_review_count", "google评论数量": "google_review_count", "google review count": "google_review_count", "google_review_count": "google_review_count",
    "yelp链接": "yelp_url", "yelp url": "yelp_url", "yelp_url": "yelp_url", "yelp评分": "yelp_rating", "yelp rating": "yelp_rating", "yelp_rating": "yelp_rating",
    "yelp评论数": "yelp_review_count", "yelp评论数量": "yelp_review_count", "yelp review count": "yelp_review_count", "yelp_review_count": "yelp_review_count",
    "社交平台": "social_profiles", "社媒": "social_profiles", "social_profiles": "social_profiles", "social profiles": "social_profiles", "近期差评": "recent_negative_reviews", "差评摘要": "recent_negative_reviews", "recent_negative_reviews": "recent_negative_reviews",
    "内容更新情况": "content_update_summary", "内容更新": "content_update_summary", "content_update_summary": "content_update_summary", "内容最近更新时间": "content_last_updated_at", "content更新时间": "content_last_updated_at", "content_last_updated_at": "content_last_updated_at",
}


def _normalize_header(value: Any) -> str:
    """Make spreadsheet headers tolerant of spaces, punctuation, and casing."""
    return re.sub(r"[\s_/\-()（）:：]+", "", str(value or "").strip().casefold())


HEADER_LOOKUP = {_normalize_header(header): field for header, field in CSV_HEADERS.items()}


def _map_import_row(row: dict[Any, Any]) -> tuple[dict[str, Any], list[str]]:
    mapped: dict[str, Any] = {}
    unknown_headers: list[str] = []
    for key, value in row.items():
        original_key = str(key or "").strip()
        if not original_key:
            continue
        normalized_key = _normalize_header(original_key)
        field = HEADER_LOOKUP.get(normalized_key)
        if not field:
            unknown_headers.append(original_key)
            continue
        if value is not None and str(value).strip() != "":
            clean_value = str(value).strip() if not isinstance(value, datetime) else value
            # Some source sheets combine two visible columns into one header.
            # Route a URL to website and the other value to phone instead of
            # shifting the rest of the row into the wrong fields.
            if normalized_key in {"电话网站", "电话官网"}:
                text = str(clean_value)
                if re.search(r"(?:https?://|www\.|[\w.-]+\.(?:com|net|org|co|us)(?:/|$))", text, re.I):
                    mapped["website"] = clean_value
                else:
                    mapped["phone"] = clean_value
            elif normalized_key in {"googlerating", "google评分"}:
                mapped["rating"] = clean_value
                mapped["google_rating"] = clean_value
            else:
                mapped[field] = clean_value
    return mapped, unknown_headers


def _parse_import_values(mapped: dict[str, Any]) -> dict[str, Any]:
    if mapped.get("website"):
        website_text = str(mapped["website"]).strip()
        url_match = re.search(r"https?://[^\s\]\)]+", website_text, re.I)
        if url_match:
            mapped["website"] = url_match.group().rstrip(".,;:")

    if mapped.get("region"):
        parts = [part.strip() for part in str(mapped.pop("region")).split(",") if part.strip()]
        if parts and re.fullmatch(r"(?:US|USA|United States|美国)", parts[-1], re.I):
            mapped.setdefault("country", "US")
            parts.pop()
        if parts and re.fullmatch(r"[A-Z]{2}", parts[-1], re.I):
            mapped.setdefault("state", parts.pop().upper())
            mapped.setdefault("country", "US")
        if parts:
            mapped.setdefault("city", ", ".join(parts))

    for field in ("rating", "google_rating", "yelp_rating"):
        if mapped.get(field) not in (None, ""):
            # Source tables often use "★4.6" or "待核验". Keep the useful
            # number and leave an unverified rating blank instead of rejecting
            # the entire merchant row.
            rating_text = str(mapped[field]).replace(",", "").strip()
            match = re.search(r"[⭐★]\s*(\d+(?:\.\d+)?)", rating_text)
            if not match:
                match = re.match(r"\s*(\d+(?:\.\d+)?)", rating_text)
            mapped[field] = float(match.group(1)) if match else None
    for field in ("google_review_count", "yelp_review_count"):
        if mapped.get(field) not in (None, ""):
            match = re.search(r"\d+(?:\.\d+)?", str(mapped[field]).replace(",", ""))
            mapped[field] = int(float(match.group())) if match else None
    return mapped


def _decode_csv(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise HTTPException(status_code=400, detail="无法识别文件编码，请另存为 UTF-8 CSV 后重试")


async def _manager_scope(db: AsyncSession, user: UserResponse):
    if _role(user) in ADMIN_ROLES:
        return None
    # A supervisor only sees pool records they imported themselves.
    return MerchantPool.created_by_id == _employee_id(user)


async def _resolve_sales_assignee(
    db: AsyncSession, assigned_sales_id: Optional[int], user: UserResponse
) -> tuple[Optional[int], Optional[str], Optional[int]]:
    """Validate assignment once so bulk and single conversion follow identical rules."""
    manager_id = _employee_id(user) if _role(user) == "sales_manager" else None
    if assigned_sales_id is None:
        return None, None, manager_id
    employee = (await db.execute(
        select(Employees).where(
            Employees.id == assigned_sales_id,
            Employees.role == "sales",
            Employees.status.in_(["active", "probation"]),
        )
    )).scalar_one_or_none()
    if not employee:
        raise HTTPException(status_code=400, detail="只能分配给在职电话销售账号")
    if _role(user) == "sales_manager" and employee.supervisor != (user.name or ""):
        raise HTTPException(status_code=403, detail="只能分配给自己的直属销售")
    return int(employee.id), employee.name, manager_id


async def _create_sales_lead_from_merchant(
    db: AsyncSession,
    merchant: MerchantPool,
    user: UserResponse,
    assigned_sales_id: Optional[int],
    assigned_sales_name: Optional[str],
    team_manager_id: Optional[int],
) -> SalesLeads:
    lead = SalesLeads(
        business_name=merchant.business_name,
        contact_name=merchant.contact_name,
        phone=merchant.phone or "",
        industry=merchant.industry,
        country=merchant.country,
        state=merchant.state,
        city=merchant.city,
        address=merchant.address,
        website=merchant.website,
        source=f"merchant_pool:{merchant.data_source}",
        merchant_pool_id=merchant.id,
        assigned_sales_id=assigned_sales_id,
        assigned_sales_name=assigned_sales_name,
        team_manager_id=team_manager_id,
        created_by_id=_employee_id(user),
        created_by_name=user.name,
    )
    latest_analysis = (await db.execute(
        select(MerchantAiAnalyses)
        .where(MerchantAiAnalyses.merchant_id == merchant.id)
        .order_by(MerchantAiAnalyses.generated_at.desc())
        .limit(1)
    )).scalar_one_or_none()
    if latest_analysis:
        lead.analysis_snapshot = latest_analysis.analysis_json
    db.add(lead)
    await db.flush()
    merchant.pool_status = "converted"
    merchant.converted_lead_id = lead.id
    merchant.isolation_reason = "已人工确认并转入独立电话销售线索库"
    return lead


async def _match_customer(db: AsyncSession, record: MerchantRecord) -> tuple[Optional[int], Optional[str]]:
    customers = (await db.execute(select(Customers))).scalars().all()
    phone = _normalize_phone(record.phone)
    name = _normalize_text(record.business_name)
    address = _normalize_text(record.address)
    website = _normalize_website(record.website)
    for customer in customers:
        if phone and phone == _normalize_phone(customer.phone):
            return customer.id, "电话已存在于正式客户"
        if website and website == _normalize_website(customer.website):
            return customer.id, "网站已存在于正式客户"
        if name and address and name == _normalize_text(customer.business_name) and address == _normalize_text(customer.address):
            return customer.id, "商家名称和地址已存在于正式客户"
    return None, None


async def _match_pool_duplicate(
    db: AsyncSession, record: MerchantRecord, exclude_id: Optional[int] = None
) -> tuple[Optional[int], Optional[str]]:
    # Historical rows still reserve the merchant identity. Otherwise an archived
    # or converted prospect can be imported again and pollute the sales pool.
    query = select(MerchantPool)
    if exclude_id is not None:
        query = query.where(MerchantPool.id != exclude_id)
    rows = (await db.execute(query)).scalars().all()
    phone = _normalize_phone(record.phone)
    name = _normalize_text(record.business_name)
    address = _normalize_text(record.address)
    website = _normalize_website(record.website)
    for row in rows:
        duplicate_id = row.duplicate_of_id or row.id
        status_hint = {
            "converted": "（该商家已转入电话销售线索）",
            ARCHIVED_STATUS: "（该商家已有归档记录）",
        }.get(row.pool_status, "")
        if phone and phone == _normalize_phone(row.phone):
            return duplicate_id, f"电话与商家池历史记录重复{status_hint}"
        if website and website == _normalize_website(row.website):
            return duplicate_id, f"网站与商家池历史记录重复{status_hint}"
        if name and address and name == _normalize_text(row.business_name) and address == _normalize_text(row.address):
            return duplicate_id, f"商家名称和地址与商家池历史记录重复{status_hint}"
    return None, None


async def _classify_record(
    db: AsyncSession, record: MerchantRecord, exclude_id: Optional[int] = None
) -> tuple[str, Optional[str], Optional[int], Optional[int]]:
    customer_id, customer_reason = await _match_customer(db, record)
    if customer_id:
        return "existing_customer", customer_reason, None, customer_id
    duplicate_id, duplicate_reason = await _match_pool_duplicate(db, record, exclude_id)
    if duplicate_id:
        return "duplicate", duplicate_reason, duplicate_id, None
    if not _normalize_phone(record.phone):
        return "no_phone", "未提供可用电话，暂不进入线索库", None, None
    if _is_closed(record.business_status):
        return "closed", "商家状态为已关闭或停业", None, None
    return "pending", None, None, None


async def _store_record(
    db: AsyncSession,
    record: MerchantRecord,
    data_source: str,
    user: UserResponse,
    raw_record: Optional[dict[str, Any]] = None,
) -> MerchantPool:
    pool_status, reason, duplicate_id, customer_id = await _classify_record(db, record)
    collected_at = record.collected_at or datetime.now(timezone.utc)
    merchant = MerchantPool(
        **record.model_dump(exclude={"collected_at"}),
        data_source=data_source.strip() or "api",
        collected_at=collected_at,
        pool_status=pool_status,
        isolation_reason=reason,
        duplicate_of_id=duplicate_id,
        existing_customer_id=customer_id,
        raw_payload=json.dumps(raw_record or record.model_dump(mode="json"), ensure_ascii=False, default=str),
        created_by_id=_employee_id(user),
        created_by_name=user.name,
    )
    db.add(merchant)
    await db.flush()
    return merchant


@router.get("/stats")
async def get_merchant_pool_stats(
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_pool_role(current_user)
    query = select(MerchantPool.pool_status, func.count(MerchantPool.id)).group_by(MerchantPool.pool_status)
    scope = await _manager_scope(db, current_user)
    if scope is not None:
        query = query.where(scope)
    counts = {status: count for status, count in (await db.execute(query)).all()}
    return {
        "total": sum(count for status, count in counts.items() if status != ARCHIVED_STATUS),
        "pending": counts.get("pending", 0),
        "isolated": sum(counts.get(item, 0) for item in ISOLATED_STATUSES),
        "converted": counts.get("converted", 0),
        "duplicates": counts.get("duplicate", 0),
        "archived": counts.get(ARCHIVED_STATUS, 0),
    }


@router.get("", response_model=MerchantListResponse)
async def list_merchant_pool(
    search: Optional[str] = None,
    pool_status: Optional[str] = None,
    region: Optional[str] = None,
    industry: Optional[str] = None,
    source: Optional[str] = None,
    rating_min: Optional[float] = Query(None, ge=0, le=5),
    skip: int = Query(0, ge=0),
    limit: int = Query(20, ge=1, le=200),
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_pool_role(current_user)
    conditions = []
    scope = await _manager_scope(db, current_user)
    if scope is not None:
        conditions.append(scope)
    if search:
        term = f"%{search.strip()}%"
        conditions.append(or_(MerchantPool.business_name.ilike(term), MerchantPool.phone.ilike(term), MerchantPool.website.ilike(term)))
    if pool_status:
        conditions.append(MerchantPool.pool_status == pool_status)
    if region:
        region_values = [value.strip() for value in region.split(",") if value.strip()]
        if len(region_values) > 1:
            conditions.append(or_(*(
                func.upper(func.trim(MerchantPool.state)) == value.upper()
                for value in region_values
            )))
        elif region_values:
            term = f"%{region_values[0]}%"
            conditions.append(or_(MerchantPool.country.ilike(term), MerchantPool.state.ilike(term), MerchantPool.city.ilike(term)))
    if industry:
        conditions.append(MerchantPool.industry.ilike(f"%{industry.strip()}%"))
    if source:
        conditions.append(MerchantPool.data_source == source)
    if rating_min is not None:
        conditions.append(MerchantPool.rating >= rating_min)
    if not pool_status:
        conditions.append(MerchantPool.pool_status != ARCHIVED_STATUS)
    query = select(MerchantPool)
    count_query = select(func.count(MerchantPool.id))
    if conditions:
        query = query.where(*conditions)
        count_query = count_query.where(*conditions)
    total = (await db.execute(count_query)).scalar() or 0
    items = (await db.execute(query.order_by(MerchantPool.id.desc()).offset(skip).limit(limit))).scalars().all()
    return {"items": items, "total": total, "skip": skip, "limit": limit}


@router.post("/import")
async def import_merchant_records(
    payload: MerchantImportRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_pool_role(current_user)
    created = []
    errors = []
    counts: dict[str, int] = {"pending": 0, "no_phone": 0, "duplicate": 0, "existing_customer": 0, "closed": 0}
    for index, raw_record in enumerate(payload.records, start=1):
        try:
            record = MerchantRecord.model_validate(_parse_import_values(dict(raw_record)))
        except (ValidationError, TypeError, ValueError) as exc:
            if isinstance(exc, ValidationError):
                detail = exc.errors()[0]
                field = ".".join(str(part) for part in detail.get("loc", ()))
                reason = f"{field}：{detail.get('msg', '格式不正确')}"
            else:
                reason = str(exc)
            errors.append({"row": index, "reason": reason})
            continue
        merchant = await _store_record(db, record, payload.data_source, current_user, raw_record=raw_record)
        created.append(merchant)
        counts[merchant.pool_status] = counts.get(merchant.pool_status, 0) + 1
    if not created:
        example = "；".join(f"第{item['row']}行 {item['reason']}" for item in errors[:3])
        raise HTTPException(status_code=400, detail=f"没有可导入的数据。{example or '请检查商家名称和数据格式'}")
    await db.commit()
    return {
        "total": len(created), "counts": counts, "errors": errors,
        "items": [MerchantResponse.model_validate(item).model_dump() for item in created],
    }


@router.post("/import-csv")
async def import_merchant_csv(
    file: UploadFile = File(...),
    data_source: str = Query("csv"),
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_pool_role(current_user)
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise HTTPException(status_code=400, detail="请上传 CSV 或 Excel（.xlsx）文件")
    content = await file.read()
    if filename.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            if not rows:
                raise HTTPException(status_code=400, detail="Excel 文件没有可导入的数据")
            headers = [str(value or "").strip() for value in rows[0]]
            reader = [dict(zip(headers, values)) for values in rows[1:] if any(value not in (None, "") for value in values)]
        except HTTPException:
            raise
        except Exception as exc:
            raise HTTPException(status_code=400, detail=f"Excel 文件读取失败：{exc}")
    else:
        text = _decode_csv(content)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    created = []
    errors = []
    recognized_fields: set[str] = set()
    ignored_headers: set[str] = set()
    counts: dict[str, int] = {"pending": 0, "no_phone": 0, "duplicate": 0, "existing_customer": 0, "closed": 0}
    for index, row in enumerate(reader, start=2):
        mapped, unknown_headers = _map_import_row(row)
        ignored_headers.update(unknown_headers)
        recognized_fields.update(mapped.keys())
        try:
            record = MerchantRecord.model_validate(_parse_import_values(mapped))
            merchant = await _store_record(db, record, data_source, current_user, raw_record=dict(row))
            created.append(merchant)
            counts[merchant.pool_status] = counts.get(merchant.pool_status, 0) + 1
        except Exception as exc:
            errors.append({"row": index, "reason": str(exc)})
    await db.commit()
    if not recognized_fields:
        raise HTTPException(status_code=400, detail="未识别到可导入列。请确认第一行包含“商家名称/名称/Business Name”等表头")
    return {
        "total": len(created), "counts": counts, "errors": errors,
        "recognized_fields": sorted(recognized_fields), "ignored_headers": sorted(ignored_headers),
    }


@router.put("/{merchant_id}", response_model=MerchantResponse)
async def update_merchant_pool_record(
    merchant_id: int,
    payload: MerchantUpdate,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_pool_role(current_user)
    query = select(MerchantPool).where(MerchantPool.id == merchant_id)
    scope = await _manager_scope(db, current_user)
    if scope is not None:
        query = query.where(scope)
    merchant = (await db.execute(query)).scalar_one_or_none()
    if not merchant:
        raise HTTPException(status_code=404, detail="商家池记录不存在或不在权限范围内")
    updates = payload.model_dump(exclude_unset=True)
    for key, value in updates.items():
        setattr(merchant, key, value)
    record = MerchantRecord.model_validate({
        "business_name": merchant.business_name, "contact_name": merchant.contact_name, "phone": merchant.phone,
        "industry": merchant.industry, "country": merchant.country, "state": merchant.state, "city": merchant.city,
        "address": merchant.address, "website": merchant.website, "rating": merchant.rating,
        "google_business_url": merchant.google_business_url, "google_rating": merchant.google_rating,
        "google_review_count": merchant.google_review_count, "yelp_url": merchant.yelp_url,
        "yelp_rating": merchant.yelp_rating, "yelp_review_count": merchant.yelp_review_count,
        "social_profiles": merchant.social_profiles, "recent_negative_reviews": merchant.recent_negative_reviews,
        "content_update_summary": merchant.content_update_summary, "content_last_updated_at": merchant.content_last_updated_at,
        "source_record_id": merchant.source_record_id, "business_status": merchant.business_status,
        "collected_at": merchant.collected_at,
    })
    merchant.pool_status, merchant.isolation_reason, merchant.duplicate_of_id, merchant.existing_customer_id = await _classify_record(
        db, record, merchant.id
    )
    merchant.updated_at = datetime.now(timezone.utc)
    await db.commit()
    await db.refresh(merchant)
    return merchant


@router.post("/{merchant_id}/archive", response_model=MerchantResponse)
async def archive_merchant_pool_record(
    merchant_id: int,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_admin_role(current_user)
    merchant = (await db.execute(select(MerchantPool).where(MerchantPool.id == merchant_id))).scalar_one_or_none()
    if not merchant:
        raise HTTPException(status_code=404, detail="商家池记录不存在")
    if merchant.pool_status == "converted" or merchant.converted_lead_id:
        raise HTTPException(status_code=409, detail="已转入电话销售线索的商家不能归档，请在线索库处理")
    if merchant.pool_status == ARCHIVED_STATUS:
        return merchant

    merchant.pool_status = ARCHIVED_STATUS
    merchant.isolation_reason = f"管理员归档：{current_user.name or current_user.email or '系统管理员'}"
    merchant.duplicate_of_id = None
    merchant.existing_customer_id = None
    merchant.updated_at = datetime.now(timezone.utc)
    await db.commit()
    await db.refresh(merchant)
    return merchant


@router.delete("/{merchant_id}")
async def delete_merchant_pool_record(
    merchant_id: int,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_admin_role(current_user)
    merchant = (await db.execute(select(MerchantPool).where(MerchantPool.id == merchant_id))).scalar_one_or_none()
    if not merchant:
        raise HTTPException(status_code=404, detail="商家池记录不存在")
    if merchant.pool_status == "converted" or merchant.converted_lead_id:
        raise HTTPException(status_code=409, detail="已转入电话销售线索的商家不能永久删除")

    await db.delete(merchant)
    await db.commit()
    return {"message": "商家池记录已删除", "id": merchant_id}


@router.post("/{merchant_id}/convert-to-lead")
async def convert_to_sales_lead(
    merchant_id: int,
    payload: ConvertToLeadRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    _ensure_pool_role(current_user)
    query = select(MerchantPool).where(MerchantPool.id == merchant_id)
    scope = await _manager_scope(db, current_user)
    if scope is not None:
        query = query.where(scope)
    merchant = (await db.execute(query)).scalar_one_or_none()
    if not merchant:
        raise HTTPException(status_code=404, detail="商家池记录不存在或不在权限范围内")
    if merchant.pool_status != "pending":
        raise HTTPException(status_code=400, detail="只有通过清洗的待处理商家可以转入销售线索")

    assigned_id, assigned_name, manager_id = await _resolve_sales_assignee(
        db, payload.assigned_sales_id, current_user
    )
    lead = await _create_sales_lead_from_merchant(
        db, merchant, current_user, assigned_id, assigned_name, manager_id
    )
    await db.commit()
    return {"message": "已转入电话销售线索库", "lead_id": lead.id, "merchant_id": merchant.id}


@router.post("/bulk-convert-to-lead")
async def bulk_convert_to_sales_leads(
    payload: BulkConvertToLeadRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Move a manager-approved batch into one salesperson's private lead library."""
    _ensure_pool_role(current_user)
    merchant_ids = list(dict.fromkeys(payload.merchant_ids))
    assigned_id, assigned_name, manager_id = await _resolve_sales_assignee(
        db, payload.assigned_sales_id, current_user
    )
    query = select(MerchantPool).where(MerchantPool.id.in_(merchant_ids))
    scope = await _manager_scope(db, current_user)
    if scope is not None:
        query = query.where(scope)
    merchants = (await db.execute(query)).scalars().all()
    merchant_by_id = {merchant.id: merchant for merchant in merchants}

    missing_ids = [merchant_id for merchant_id in merchant_ids if merchant_id not in merchant_by_id]
    invalid = [merchant for merchant in merchants if merchant.pool_status != "pending"]
    if missing_ids:
        raise HTTPException(status_code=404, detail=f"有 {len(missing_ids)} 条商家记录不存在或不在您的权限范围内")
    if invalid:
        raise HTTPException(
            status_code=400,
            detail=f"有 {len(invalid)} 条记录不是待清洗可用状态，无法重复分配",
        )

    leads = []
    for merchant_id in merchant_ids:
        leads.append(await _create_sales_lead_from_merchant(
            db, merchant_by_id[merchant_id], current_user, assigned_id, assigned_name, manager_id
        ))
    await db.commit()
    return {
        "message": f"已将 {len(leads)} 条可用商家分配给 {assigned_name}",
        "assigned_sales_id": assigned_id,
        "assigned_sales_name": assigned_name,
        "converted_count": len(leads),
        "lead_ids": [lead.id for lead in leads],
    }


@router.post("/bulk-update-industry")
async def bulk_update_merchant_industry(
    payload: BulkIndustryUpdateRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Apply one reviewed industry label to selected non-final pool records."""
    _ensure_pool_role(current_user)
    merchant_ids = list(dict.fromkeys(payload.merchant_ids))
    query = select(MerchantPool).where(MerchantPool.id.in_(merchant_ids))
    scope = await _manager_scope(db, current_user)
    if scope is not None:
        query = query.where(scope)
    merchants = (await db.execute(query)).scalars().all()
    merchant_by_id = {merchant.id: merchant for merchant in merchants}
    missing_ids = [merchant_id for merchant_id in merchant_ids if merchant_id not in merchant_by_id]
    if missing_ids:
        raise HTTPException(status_code=404, detail=f"有 {len(missing_ids)} 条记录不存在或不在您的权限范围内")

    locked = [merchant for merchant in merchants if merchant.pool_status in {"converted", ARCHIVED_STATUS}]
    if locked:
        raise HTTPException(status_code=409, detail="已转线索或已归档的记录不能批量修改行业")

    updated_at = datetime.now(timezone.utc)
    for merchant in merchants:
        merchant.industry = payload.industry
        merchant.updated_at = updated_at
    await db.commit()
    return {"message": f"已更新 {len(merchants)} 条商家的行业", "updated_count": len(merchants), "industry": payload.industry}


@router.post("/bulk-delete")
async def bulk_delete_merchant_pool_records(
    payload: BulkMerchantIdsRequest,
    current_user: UserResponse = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Permanently remove selected pool records without touching leads or customers."""
    _ensure_admin_role(current_user)
    merchant_ids = list(dict.fromkeys(payload.merchant_ids))
    merchants = (await db.execute(
        select(MerchantPool).where(MerchantPool.id.in_(merchant_ids))
    )).scalars().all()
    merchant_by_id = {merchant.id: merchant for merchant in merchants}
    missing_ids = [merchant_id for merchant_id in merchant_ids if merchant_id not in merchant_by_id]
    if missing_ids:
        raise HTTPException(status_code=404, detail=f"有 {len(missing_ids)} 条商家记录不存在")

    locked = [merchant for merchant in merchants if merchant.pool_status == "converted" or merchant.converted_lead_id]
    if locked:
        raise HTTPException(status_code=409, detail="已转入电话销售线索的记录不能批量永久删除")

    for merchant in merchants:
        await db.delete(merchant)
    await db.commit()
    return {"message": f"已删除 {len(merchants)} 条商家池记录", "deleted_count": len(merchants)}
