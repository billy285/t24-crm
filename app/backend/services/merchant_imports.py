"""Fixed-template import receipts, validation and safe business-record creation."""

import csv
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone
from pathlib import PurePath
from typing import Any
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.exc import IntegrityError

from models.customers import Customers
from models.merchant_ai_analyses import MerchantAiAnalyses
from models.merchant_import_batches import MerchantImportBatch
from models.merchant_pool import MerchantPool
from models.sales_leads import SalesLeads
from routers.merchant_pool import (
    ADMIN_ROLES, IMPORT_TEMPLATE_FIELDS, IMPORT_TEMPLATE_HEADERS, MerchantRecord,
    _decode_csv, _employee_id, _ensure_pool_role, _is_closed, _normalize_text,
    _normalize_website, _parse_import_values, _role, _validate_import_template_headers,
)
from schemas.auth import UserResponse
from services.phone_numbers import parse_phone_number, phone_match_key
from services.merchant_row_locks import lock_merchant_rows

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 30 * 1024 * 1024
TEMPLATE_VERSION = "fixed-five-v1"


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, default=str, sort_keys=True)


def parse_import_file(filename: str, content: bytes) -> list[dict]:
    if not content or len(content) > MAX_FILE_BYTES:
        raise HTTPException(400, "文件为空或超过 10MB，请拆分后上传")
    filename = filename.lower()
    if filename.endswith(".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > MAX_EXPANDED_BYTES:
                    raise HTTPException(400, "Excel 展开内容过大，请使用精简模板")
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            try:
                iterator = workbook.active.iter_rows(values_only=True)
                headers = next(iterator, ())
                _validate_import_template_headers(list(headers))
                rows = []
                for number, values in enumerate(iterator, start=2):
                    if not any(value is not None and str(value).strip() for value in values):
                        continue
                    rows.append({"row": number, "values": [str(value) if value is not None else "" for value in values]})
                    if len(rows) > 2000:
                        raise HTTPException(400, "单次最多导入 2000 条商家记录")
            finally:
                workbook.close()
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(400, "Excel 文件无法读取，请另存为 .xlsx 后重试")
    elif filename.endswith(".csv"):
        try:
            text = _decode_csv(content)
            try:
                dialect = csv.Sniffer().sniff(text[:4096], delimiters=",\t;|")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect, strict=True)
            _validate_import_template_headers(next(reader, []))
            rows = []
            while True:
                number = reader.line_num + 1
                try:
                    values = next(reader)
                except StopIteration:
                    break
                if not any(value.strip() for value in values):
                    continue
                rows.append({"row": number, "values": values})
                if len(rows) > 2000:
                    raise HTTPException(400, "单次最多导入 2000 条商家记录")
        except csv.Error:
            raise HTTPException(400, "CSV 引号或分隔格式有误，请使用下载的模板")
    else:
        raise HTTPException(400, "请上传 CSV 或 Excel（.xlsx）文件")
    if not rows:
        raise HTTPException(400, "模板中没有商家记录，请填写后重试")
    return rows


def _identity_keys(record) -> list[str]:
    keys = []
    phone = phone_match_key(record.phone, record.country)
    if phone:
        keys.append(f"phone:{phone}")
    website = _normalize_website(record.website)
    if website:
        keys.append(f"website:{website}")
    name, address = _normalize_text(record.business_name), _normalize_text(record.address)
    if name and address:
        keys.append(f"location:{name}:{address}")
    return keys


async def evaluate_rows(db: AsyncSession, rows: list[dict]) -> list[dict]:
    # Take one identity snapshot instead of scanning every table per uploaded row.
    identities: dict[str, tuple[str, str, int | None, int | None, int | None]] = {}
    customers = (await db.execute(select(Customers))).scalars().all()
    merchants = (await db.execute(select(MerchantPool))).scalars().all()
    leads = (await db.execute(select(SalesLeads))).scalars().all()
    merchant_ids = {merchant.id for merchant in merchants}
    for records, kind in ((customers, "customer"), (merchants, "merchant"), (leads, "lead")):
        for record in records:
            for key in _identity_keys(record):
                if kind == "customer":
                    match = ("existing_customer", "与正式客户资料重复，暂不进入销售队列", None, record.id, None)
                elif kind == "merchant":
                    reference_id = record.duplicate_of_id if record.duplicate_of_id in merchant_ids else record.id
                    match = ("duplicate", "与商家池历史记录重复，需人工核对", reference_id, None, None)
                else:
                    match = ("duplicate", "与已有销售线索重复，需人工核对", None, None, None)
                identities.setdefault(key, match)
    result = []
    for row in rows:
        values = row["values"]
        raw = dict(zip(IMPORT_TEMPLATE_HEADERS, values))
        entry = {"row": row["row"], "raw": raw, "business_name": str(raw.get("商家名称") or ""), "phone": str(raw.get("商家电话") or ""), "status": "error", "reason": "", "warnings": []}
        result.append(entry)
        if len(values) != 5:
            entry["reason"] = "必须正好填写 5 列"
            continue
        mapped = {field: str(value).strip() for field, value in zip(IMPORT_TEMPLATE_FIELDS, values) if str(value).strip()}
        source = mapped.pop("data_source", "csv")
        try:
            if len(source) > 50:
                raise ValueError("来源最多填写 50 个字符")
            record = MerchantRecord.model_validate(_parse_import_values(mapped))
            if not record.business_name.strip():
                raise ValueError("商家名称不能为空")
        except (ValidationError, ValueError, TypeError) as error:
            if isinstance(error, ValidationError):
                field = error.errors()[0].get("loc", ("business_name",))[0]
                entry["reason"] = "商家名称必填且最多 200 个字符" if field == "business_name" else "字段格式不正确"
            else:
                entry["reason"] = str(error)
            continue
        phone = parse_phone_number(record.phone, record.country)
        keys = _identity_keys(record)
        duplicate = next((identities[key] for key in keys if key in identities), None)
        duplicate_row = None
        if duplicate:
            status, reason, duplicate_id, customer_id, duplicate_row = duplicate
        elif not phone.is_valid:
            status, reason, duplicate_id, customer_id = "no_phone", phone.reason or "电话需补齐后才能分配", None, None
        elif _is_closed(record.business_status):
            status, reason, duplicate_id, customer_id = "closed", "商家已停业", None, None
        else:
            status, reason, duplicate_id, customer_id = "pending", "待人工核对；尚未分配销售", None, None
        if not record.country:
            entry["warnings"].append("地区未明确国家，未默认填写美国")
        if not raw.get("来源", "").strip():
            entry["warnings"].append("来源为空，记录为 CSV 导入")
        entry.update({"status": status, "reason": reason, "data": record.model_dump(mode="json"), "data_source": source, "phone_status": phone.status, "phone_country": phone.country, "normalized_phone": phone.e164, "duplicate_of_id": duplicate_id, "existing_customer_id": customer_id, "duplicate_row": duplicate_row})
        for key in keys:
            identities.setdefault(key, ("duplicate", f"与本文件第 {row['row']} 行重复，需人工核对", None, None, row["row"]))
    return result


def summarize(rows: list[dict]) -> dict:
    counts = {"pending": 0, "isolated": 0, "errors": 0, "importable": 0}
    for row in rows:
        if row["status"] == "error":
            counts["errors"] += 1
        else:
            counts["importable"] += 1
            counts["pending" if row["status"] == "pending" else "isolated"] += 1
    return counts


def batch_response(batch: MerchantImportBatch, *, duplicate_upload: bool = False) -> dict:
    result = json.loads(batch.result_json)
    return {"id": batch.id, "filename": batch.filename, "file_sha256": batch.file_sha256, "template_version": batch.template_version, "status": batch.status, "row_count": batch.row_count, "created_at": batch.created_at, "committed_at": batch.committed_at, "reverted_at": batch.reverted_at, "created_by_name": batch.created_by_name, "duplicate_upload": duplicate_upload, "summary": summarize(result), "rows": [{key: value for key, value in row.items() if key not in {"snapshot", "data", "duplicate_of_id", "existing_customer_id"}} for row in result]}


def batch_scope(user: UserResponse):
    _ensure_pool_role(user)
    return None if _role(user) in ADMIN_ROLES else MerchantImportBatch.created_by_id == _employee_id(user)


async def get_batch(db: AsyncSession, user: UserResponse, batch_id: str) -> MerchantImportBatch:
    scope = batch_scope(user)
    query = select(MerchantImportBatch).where(MerchantImportBatch.id == batch_id)
    if scope is not None:
        query = query.where(scope)
    batch = (await db.execute(query)).scalar_one_or_none()
    if not batch:
        raise HTTPException(404, "导入批次不存在或不在您的权限范围内")
    return batch


async def preview_import(db: AsyncSession, user: UserResponse, filename: str, content: bytes) -> dict:
    scope = batch_scope(user)
    rows = parse_import_file(filename, content)
    digest = hashlib.sha256(content).hexdigest()
    query = select(MerchantImportBatch).where(MerchantImportBatch.file_sha256 == digest, MerchantImportBatch.status.in_(["preview", "committed"]))
    if scope is not None:
        query = query.where(scope)
    previous = (await db.execute(query.order_by(MerchantImportBatch.created_at.desc()).limit(1))).scalar_one_or_none()
    if previous:
        # Do not rewrite an existing receipt: a concurrent confirmation may have
        # already recorded merchant IDs and undo snapshots. Confirm revalidates.
        return batch_response(previous, duplicate_upload=True)
    checked = await evaluate_rows(db, rows)
    batch = MerchantImportBatch(id=str(uuid4()), filename=PurePath(filename).name[:255], file_sha256=digest, active_file_key=f"{_employee_id(user)}:{digest}", template_version=TEMPLATE_VERSION, created_by_id=_employee_id(user), created_by_name=user.name, status="preview", row_count=len(rows), rows_json=_json(rows), result_json=_json(checked))
    db.add(batch)
    try:
        await db.commit()
    except IntegrityError:
        await db.rollback()
        previous = (await db.execute(select(MerchantImportBatch).where(MerchantImportBatch.active_file_key == f"{_employee_id(user)}:{digest}"))).scalar_one_or_none()
        if not previous:
            raise
        return batch_response(previous, duplicate_upload=True)
    await db.refresh(batch)
    return batch_response(batch)


def merchant_snapshot(merchant: MerchantPool) -> str:
    # Include every persisted merchant field, including update time. A later edit,
    # enrichment or conversion means this row must no longer be undone by import.
    return _json({column.name: getattr(merchant, column.name) for column in MerchantPool.__table__.columns})


async def confirm_import(db: AsyncSession, user: UserResponse, batch_id: str) -> dict:
    batch_scope(user)
    await lock_merchant_rows(db, [])
    batch = await get_batch(db, user, batch_id)
    if batch.status == "committed":
        return batch_response(batch)
    if batch.status != "preview":
        raise HTTPException(409, "该批次已撤销或正在处理，请刷新后查看")
    claimed = await db.execute(update(MerchantImportBatch).where(MerchantImportBatch.id == batch.id, MerchantImportBatch.status == "preview").values(status="committing"))
    if claimed.rowcount != 1:
        raise HTTPException(409, "批次正在处理，请刷新后查看结果")
    # The preview is informative; current database identities are authoritative.
    checked = await evaluate_rows(db, json.loads(batch.rows_json))
    reference_ids = {row["duplicate_of_id"] for row in checked if row.get("duplicate_of_id")}
    if reference_ids:
        # Serialize new references with batch undo. A reference can disappear
        # while we wait for the lock, so re-evaluate against the current rows.
        references = await lock_merchant_rows(db, reference_ids)
        locked_ids = {merchant.id for merchant in references}
        checked = await evaluate_rows(db, json.loads(batch.rows_json))
        current_references = {row["duplicate_of_id"] for row in checked if row.get("duplicate_of_id")}
        if not current_references.issubset(locked_ids):
            await db.rollback()
            raise HTTPException(409, "商家资料刚被其他操作更新，批次尚未入池，请重新确认")
    if not summarize(checked)["importable"]:
        await db.rollback()
        raise HTTPException(400, "没有可入池的记录，请下载错误行并修正文件")
    now = datetime.now(timezone.utc)
    created_by_row = {}
    for row in checked:
        if row["status"] == "error":
            continue
        record = MerchantRecord.model_validate(row["data"])
        duplicate_id = row["duplicate_of_id"] or created_by_row.get(row.get("duplicate_row"))
        merchant = MerchantPool(**record.model_dump(exclude={"collected_at"}), data_source=row["data_source"], collected_at=record.collected_at or now, pool_status=row["status"], isolation_reason=None if row["status"] == "pending" else row["reason"], duplicate_of_id=duplicate_id, existing_customer_id=row["existing_customer_id"], raw_payload=_json(row["raw"]), created_by_id=batch.created_by_id, created_by_name=batch.created_by_name)
        db.add(merchant)
        await db.flush()
        # Refresh SQLite's UTC timezone-less timestamps before storing the receipt.
        await db.refresh(merchant)
        created_by_row[row["row"]] = merchant.id
        row.update({"merchant_id": merchant.id, "snapshot": merchant_snapshot(merchant)})
    batch.status = "committed"
    batch.committed_at = now
    batch.result_json = _json(checked)
    await db.commit()
    return batch_response(batch)


async def revert_import(db: AsyncSession, user: UserResponse, batch_id: str) -> dict:
    batch_scope(user)
    await lock_merchant_rows(db, [])
    batch = await get_batch(db, user, batch_id)
    if batch.status == "reverted":
        return batch_response(batch)
    if batch.status != "committed":
        raise HTTPException(409, "只有已入池批次可以撤销")
    claimed = await db.execute(update(MerchantImportBatch).where(MerchantImportBatch.id == batch.id, MerchantImportBatch.status == "committed").values(status="reverting"))
    if claimed.rowcount != 1:
        raise HTTPException(409, "批次正在处理，请刷新后重试")
    rows = json.loads(batch.result_json)
    merchant_ids = [row["merchant_id"] for row in rows if row.get("merchant_id")]
    merchants = await lock_merchant_rows(db, merchant_ids)
    by_id = {merchant.id: merchant for merchant in merchants}
    linked_leads = set((await db.execute(select(SalesLeads.merchant_pool_id).where(SalesLeads.merchant_pool_id.in_(merchant_ids)))).scalars().all())
    analyzed_ids = set((await db.execute(select(MerchantAiAnalyses.merchant_id).where(MerchantAiAnalyses.merchant_id.in_(merchant_ids)))).scalars().all())
    referenced_ids = set((await db.execute(select(MerchantPool.duplicate_of_id).where(MerchantPool.duplicate_of_id.in_(merchant_ids), MerchantPool.id.notin_(merchant_ids)))).scalars().all())
    conflicts = []
    for row in rows:
        merchant_id = row.get("merchant_id")
        if not merchant_id:
            continue
        merchant = by_id.get(merchant_id)
        if not merchant or merchant.converted_lead_id or merchant.pool_status == "converted" or merchant_id in linked_leads or merchant_id in analyzed_ids or merchant_id in referenced_ids or merchant_snapshot(merchant) != row.get("snapshot"):
            conflicts.append(row["row"])
    if conflicts:
        await db.rollback()
        raise HTTPException(409, {"message": "批次已有记录被修改、分析、引用、删除或转为线索，整批未撤销", "rows": conflicts})
    for merchant in merchants:
        await db.delete(merchant)
    batch.status = "reverted"
    batch.active_file_key = None
    batch.reverted_at = datetime.now(timezone.utc)
    await db.commit()
    return batch_response(batch)


def error_csv(batch: MerchantImportBatch) -> str:
    output = io.StringIO(newline="")
    writer = csv.writer(output)
    writer.writerow(["原始行号", *IMPORT_TEMPLATE_HEADERS, "处理结果", "原因"])
    for row in json.loads(batch.result_json):
        if row["status"] == "pending":
            continue
        # CSV cells are data, never formulas when opened in Excel.
        values = [row["row"], *(row["raw"].get(header, "") for header in IMPORT_TEMPLATE_HEADERS), row["status"], row["reason"]]
        writer.writerow([f"'{value}" if str(value).lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else value for value in values])
    return "\ufeff" + output.getvalue()
